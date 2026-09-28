"""
Client Report Export Module
Generates comprehensive white-label client deliverables:
1. Multi-sheet Excel workbook (report.xlsx) via openpyxl:
   - Summary, Link targets, Competitor gaps, Brand position, Pitches, Content brief
2. Self-contained, executive HTML client report (report.html):
   - Agency white-label header, KPI cards, Brand visibility, Top targets, Competitor gaps, Linkable asset ideas
"""

import io
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_excel_sheet(ws, header_fill_color="1E293B", header_font_color="FFFFFF"):
    """Apply professional formatting, header colors, borders, and auto-fit column widths."""
    header_font = Font(name="Segoe UI", size=11, bold=True, color=header_font_color)
    header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)


def export_excel_report(
    run_payload: dict,
    table_df: pd.DataFrame,
    gap_pages_df: pd.DataFrame,
    brand_summary_df: pd.DataFrame,
    pitches_list: list[dict],
    content_brief_data: dict,
    agency_name: str = "Digital PR & SEO Intelligence"
) -> bytes:
    """
    Generate professional multi-sheet Excel report (report.xlsx) using openpyxl.
    """
    client_dict = run_payload.get("client", {})
    client_name = client_dict.get("name", "Client")
    client_domain = client_dict.get("domain", "")
    service = run_payload.get("service", "SEO Link Building")
    location = run_payload.get("location", "Global")
    created = run_payload.get("created", datetime.datetime.now().isoformat())
    date_str = str(created)[:10]

    output = io.BytesIO()
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Sheet: Summary
    ws_summary = wb.create_sheet(title="Summary")
    summary_data = [
        ["Report Property", "Value"],
        ["Client Brand Name", client_name],
        ["Client Domain", client_domain],
        ["Service / Niche", service],
        ["Target Market(s)", location],
        ["Analysis Date", date_str],
        ["White-Label Prepared By", agency_name],
        ["Total AI Answers Analyzed", len(run_payload.get("records", []))],
        ["Total Unique Domains Cited", len(table_df) if not table_df.empty else 0],
        ["Priority Outreach Targets", int(table_df["action"].str.startswith("Outreach").sum()) if not table_df.empty else 0],
        ["Competitor Gap Pages Found", len(gap_pages_df) if not gap_pages_df.empty else 0],
        ["Gemini Model Used", run_payload.get("model", "gemini-3.1-flash-lite")],
        ["Total Gemini Tokens Consumed", run_payload.get("total_tokens", 0)],
    ]
    for r in summary_data:
        ws_summary.append(r)
    format_excel_sheet(ws_summary, header_fill_color="1E88E5")

    # 2. Sheet: Link targets
    ws_targets = wb.create_sheet(title="Link targets")
    if not table_df.empty:
        export_cols = [
            c for c in [
                "domain", "priority_score", "citations", "prompts_cited_in",
                "competitor_gap_pages", "best_pitch_type", "citation_worthiness_score",
                "contact", "guest_post_url", "action", "newest_last_updated", "top_urls"
            ] if c in table_df.columns
        ]
        df_t = table_df[export_cols].copy()
        # Rename for clean export
        df_t.columns = [c.replace("_", " ").title() for c in df_t.columns]
        ws_targets.append(list(df_t.columns))
        for row in df_t.itertuples(index=False):
            ws_targets.append(list(row))
    else:
        ws_targets.append(["No link target data available"])
    format_excel_sheet(ws_targets, header_fill_color="0F172A")

    # 3. Sheet: Competitor gaps
    ws_gaps = wb.create_sheet(title="Competitor gaps")
    if not gap_pages_df.empty:
        gap_cols = [
            c for c in [
                "url", "domain", "page_title", "competitors_present", "pitch_type",
                "citation_worthiness_score", "contact", "last_updated", "word_count", "fetch_status"
            ] if c in gap_pages_df.columns
        ]
        df_g = gap_pages_df[gap_cols].copy()
        if "competitors_present" in df_g.columns:
            df_g["competitors_present"] = df_g["competitors_present"].apply(lambda x: ", ".join(x) if isinstance(x, list) else str(x))
        df_g.columns = [c.replace("_", " ").title() for c in df_g.columns]
        ws_gaps.append(list(df_g.columns))
        for row in df_g.itertuples(index=False):
            ws_gaps.append(list(row))
    else:
        ws_gaps.append(["No competitor gap pages found"])
    format_excel_sheet(ws_gaps, header_fill_color="DC2626")

    # 4. Sheet: Brand position
    ws_brand = wb.create_sheet(title="Brand position")
    if not brand_summary_df.empty:
        b_cols = [
            c for c in [
                "brand", "role", "mentions", "mention_rate_pct", "avg_position_display",
                "pct_top_3", "pct_positive", "pct_neutral", "pct_negative", "dominant_sentiment"
            ] if c in brand_summary_df.columns
        ]
        df_b = brand_summary_df[b_cols].copy()
        df_b.columns = [c.replace("_", " ").title() for c in df_b.columns]
        ws_brand.append(list(df_b.columns))
        for row in df_b.itertuples(index=False):
            ws_brand.append(list(row))
    else:
        ws_brand.append(["No brand position analysis available"])
    format_excel_sheet(ws_brand, header_fill_color="D97706")

    # 5. Sheet: Pitches
    ws_pitches = wb.create_sheet(title="Pitches")
    if pitches_list:
        p_rows = []
        for p in pitches_list:
            topics = "; ".join(p.get("suggested_topics", [])) if p.get("suggested_topics") else ""
            p_rows.append([
                p.get("domain", ""),
                p.get("url", ""),
                p.get("contact", ""),
                p.get("pitch_type", ""),
                p.get("subject", ""),
                p.get("email", ""),
                p.get("suggested_anchor", ""),
                topics
            ])
        ws_pitches.append(["Domain", "Target URL", "Contact", "Pitch Strategy", "Subject Line", "Email Body", "Suggested Anchor", "Guest Post Topics"])
        for r in p_rows:
            ws_pitches.append(r)
    else:
        ws_pitches.append(["Domain", "Target URL", "Contact", "Pitch Strategy", "Subject Line", "Email Body", "Suggested Anchor", "Guest Post Topics"])
        ws_pitches.append(["No pitches generated yet", "", "", "", "", "", "", ""])
    format_excel_sheet(ws_pitches, header_fill_color="2563EB")

    # 6. Sheet: Content brief
    ws_brief = wb.create_sheet(title="Content brief")
    ws_brief.append(["Section", "Content Brief Details"])
    if content_brief_data:
        # Common angles
        for a in content_brief_data.get("common_angles", []):
            ws_brief.append(["Winning Angle", str(a)])
        # Questions answered
        for q in content_brief_data.get("questions_answered", []):
            ws_brief.append(["Target Question", str(q)])
        # Data points
        for d in content_brief_data.get("data_points_used", []):
            ws_brief.append(["Benchmark Data", str(d)])
        # Format
        fmt = content_brief_data.get("recommended_format", {})
        if isinstance(fmt, dict):
            ws_brief.append(["Recommended Format", fmt.get("content_type", "")])
            ws_brief.append(["Word Count", fmt.get("recommended_word_count", "")])
            ws_brief.append(["Visual Assets", fmt.get("visual_assets", "")])
        # Linkable asset ideas
        for idx, idea in enumerate(content_brief_data.get("content_ideas", [])):
            if isinstance(idea, dict):
                ws_brief.append([f"Asset Idea #{idx+1} Title", idea.get("title", "")])
                ws_brief.append([f"Asset Idea #{idx+1} Format", idea.get("format", "")])
                ws_brief.append([f"Asset Idea #{idx+1} Why AI Cites It", idea.get("why_ai_cites_it", "")])
                ws_brief.append([f"Asset Idea #{idx+1} Outreach Angle", idea.get("target_outreach_angle", "")])
    else:
        ws_brief.append(["Notice", "No content brief built yet"])
    format_excel_sheet(ws_brief, header_fill_color="059669")

    wb.save(output)
    return output.getvalue()


def generate_html_report(
    run_payload: dict,
    table_df: pd.DataFrame,
    gap_pages_df: pd.DataFrame,
    brand_summary_df: pd.DataFrame,
    content_brief_data: dict,
    agency_name: str = "Digital PR & SEO Intelligence"
) -> str:
    """
    Generate an executive white-label single-page HTML client report (report.html).
    """
    client_dict = run_payload.get("client", {})
    client_name = client_dict.get("name", "Client")
    client_domain = client_dict.get("domain", "")
    service = run_payload.get("service", "SEO Link Building")
    location = run_payload.get("location", "the UK")
    created = run_payload.get("created", datetime.datetime.now().isoformat())
    date_str = str(created)[:10]

    total_answers = len(run_payload.get("records", []))
    total_domains = len(table_df) if not table_df.empty else 0
    total_outreach = int(table_df["action"].str.startswith("Outreach").sum()) if not table_df.empty else 0
    gap_count = len(gap_pages_df) if not gap_pages_df.empty else 0

    # Client SoV & Avg Position
    client_sov = "0%"
    client_avg_pos = "-"
    if not brand_summary_df.empty:
        c_row = brand_summary_df[brand_summary_df["brand"].str.lower() == client_name.lower().strip()]
        if not c_row.empty:
            client_sov = f"{c_row['mention_rate_pct'].iloc[0]}%"
            client_avg_pos = c_row['avg_position_display'].iloc[0]

    # Top 10 Link Targets Table Rows
    top10_targets_html = ""
    if not table_df.empty:
        top10 = table_df.head(10)
        for _, r in top10.iterrows():
            top10_targets_html += f"""
            <tr>
                <td><strong>{r.get('domain', '')}</strong></td>
                <td><span class="badge badge-primary">{r.get('priority_score', 0):.1f}</span></td>
                <td>{r.get('citations', 0)}</td>
                <td>{r.get('best_pitch_type', 'Niche edit')}</td>
                <td><span class="score-pill">{r.get('citation_worthiness_score', 50)}/100</span></td>
                <td><code>{r.get('contact', 'Contact form')}</code></td>
            </tr>
            """
    else:
        top10_targets_html = "<tr><td colspan='6'>No targets available</td></tr>"

    # Top 5 Competitor Gaps Rows
    top5_gaps_html = ""
    if not gap_pages_df.empty:
        top5_gaps = gap_pages_df.head(5)
        for _, r in top5_gaps.iterrows():
            comps = r.get("competitors_present", [])
            comp_str = ", ".join(comps) if isinstance(comps, list) else str(comps)
            top5_gaps_html += f"""
            <tr>
                <td><a href="{r.get('url', '#')}" target="_blank" style="color:#2563eb; text-decoration:none;">{r.get('page_title', r.get('url', ''))[:60]}...</a></td>
                <td><strong>{r.get('domain', '')}</strong></td>
                <td><span class="badge badge-danger">{comp_str}</span></td>
                <td>{r.get('pitch_type', 'List inclusion')}</td>
                <td><code>{r.get('contact', '-')}</code></td>
            </tr>
            """
    else:
        top5_gaps_html = "<tr><td colspan='5'>No competitor gaps detected</td></tr>"

    # 5 Linkable Asset Ideas from Content Brief
    ideas_html = ""
    if content_brief_data and "content_ideas" in content_brief_data:
        ideas = content_brief_data.get("content_ideas", [])[:5]
        for idx, idea in enumerate(ideas):
            ideas_html += f"""
            <div class="card idea-card">
                <div style="display:flex; justify-content:space-between; align-items:flex-start;">
                    <h4 style="margin:0px 0px 6px 0px; color:#0F172A; font-size:1.05rem;">
                        <span class="num-circle">{idx+1}</span> {idea.get('title', '')}
                    </h4>
                    <span class="badge badge-info">{idea.get('format', 'Data Study')}</span>
                </div>
                <p style="margin:6px 0px 4px 0px; font-size:0.9rem; color:#475569;">
                    <strong>Why AI Cites It:</strong> {idea.get('why_ai_cites_it', '')}
                </p>
                <p style="margin:0px; font-size:0.9rem; color:#2563EB;">
                    <strong>Outreach Pitch Strategy:</strong> {idea.get('target_outreach_angle', '')}
                </p>
            </div>
            """
    else:
        ideas_html = "<p style='color:#64748B;'>Click 'Build content brief' in the application to generate linkable asset strategies.</p>"

    # Brand summary table rows
    brand_rows_html = ""
    if not brand_summary_df.empty:
        for _, r in brand_summary_df.iterrows():
            brand_rows_html += f"""
            <tr>
                <td><strong>{r.get('brand', '')}</strong></td>
                <td>{r.get('role', '')}</td>
                <td><strong>{r.get('mention_rate_pct', 0)}%</strong></td>
                <td><span class="badge badge-primary">{r.get('avg_position_display', '-')}</span></td>
                <td>{r.get('pct_top_3', 0)}%</td>
                <td>{r.get('dominant_sentiment', 'Neutral')}</td>
            </tr>
            """

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AI Citation Link Intelligence Report — {client_name}</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #1E293B;
            background-color: #F8FAFC;
            margin: 0;
            padding: 24px;
            line-height: 1.5;
        }}
        .container {{
            max-width: 1100px;
            margin: 0 auto;
            background: #FFFFFF;
            border-radius: 14px;
            box-shadow: 0 4px 20px rgba(0,0,0,0.06);
            padding: 36px 44px;
        }}
        .header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            border-bottom: 2px solid #E2E8F0;
            padding-bottom: 20px;
            margin-bottom: 28px;
        }}
        .agency-title {{
            font-size: 1.25rem;
            font-weight: 700;
            color: #2563EB;
            letter-spacing: -0.5px;
        }}
        .report-title {{
            font-size: 1.8rem;
            font-weight: 800;
            color: #0F172A;
            margin: 4px 0 0 0;
        }}
        .meta-grid {{
            display: grid;
            grid-template-columns: repeat(4, 1fr);
            gap: 16px;
            margin-bottom: 28px;
        }}
        .kpi-card {{
            background: #F1F5F9;
            border-radius: 10px;
            padding: 16px;
            border-left: 4px solid #2563EB;
        }}
        .kpi-val {{
            font-size: 1.6rem;
            font-weight: 800;
            color: #0F172A;
            margin-top: 4px;
        }}
        .kpi-lbl {{
            font-size: 0.82rem;
            color: #64748B;
            text-transform: uppercase;
            font-weight: 600;
        }}
        h3 {{
            font-size: 1.25rem;
            font-weight: 700;
            color: #0F172A;
            margin-top: 32px;
            margin-bottom: 14px;
            display: flex;
            align-items: center;
            gap: 8px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 0.92rem;
            margin-bottom: 24px;
        }}
        th {{
            background: #0F172A;
            color: #FFFFFF;
            text-align: left;
            padding: 10px 14px;
            font-weight: 600;
        }}
        td {{
            padding: 10px 14px;
            border-bottom: 1px solid #E2E8F0;
        }}
        tr:nth-child(even) {{
            background-color: #F8FAFC;
        }}
        .badge {{
            display: inline-block;
            padding: 3px 8px;
            border-radius: 6px;
            font-size: 0.78rem;
            font-weight: 600;
        }}
        .badge-primary {{ background: #DBEAFE; color: #1D4ED8; }}
        .badge-danger {{ background: #FEE2E2; color: #DC2626; }}
        .badge-info {{ background: #E0E7FF; color: #4338CA; }}
        .score-pill {{
            background: #DCFCE7;
            color: #15803D;
            padding: 3px 8px;
            border-radius: 12px;
            font-weight: 700;
            font-size: 0.82rem;
        }}
        .idea-card {{
            background: #FFFFFF;
            border: 1px solid #E2E8F0;
            border-radius: 10px;
            padding: 16px 20px;
            margin-bottom: 14px;
            box-shadow: 0 2px 6px rgba(0,0,0,0.02);
        }}
        .num-circle {{
            display: inline-flex;
            width: 22px;
            height: 22px;
            background: #2563EB;
            color: #fff;
            border-radius: 50%;
            font-size: 0.75rem;
            align-items: center;
            justify-content: center;
            margin-right: 6px;
        }}
        .footer {{
            margin-top: 40px;
            padding-top: 18px;
            border-top: 1px solid #E2E8F0;
            text-align: center;
            font-size: 0.85rem;
            color: #94A3B8;
        }}
        @media print {{
            body {{ background: #FFFFFF; padding: 0; }}
            .container {{ box-shadow: none; padding: 0; }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <div>
                <div class="agency-title">{agency_name}</div>
                <div class="report-title">AI Search Grounding & Link Intelligence Report</div>
            </div>
            <div style="text-align:right;">
                <div style="font-weight:700; color:#0F172A;">Client: {client_name}</div>
                <div style="font-size:0.88rem; color:#64748B;">{client_domain}</div>
                <div style="font-size:0.82rem; color:#94A3B8;">Date: {date_str}</div>
            </div>
        </div>

        <div class="meta-grid">
            <div class="kpi-card">
                <div class="kpi-lbl">Client Share of Voice</div>
                <div class="kpi-val" style="color:#2563EB;">{client_sov}</div>
            </div>
            <div class="kpi-card">
                <div class="kpi-lbl">Client Avg AI Position</div>
                <div class="kpi-val" style="color:#059669;">{client_avg_pos}</div>
            </div>
            <div class="kpi-card">
                <div class="kpi-lbl">Outreach Targets</div>
                <div class="kpi-val">{total_outreach}</div>
            </div>
            <div class="kpi-card">
                <div class="kpi-lbl">Competitor Gaps</div>
                <div class="kpi-val" style="color:#DC2626;">{gap_count}</div>
            </div>
        </div>

        <h3>👑 1. AI Recommendation Hierarchy & Brand Visibility</h3>
        <table>
            <thead>
                <tr>
                    <th>Brand Name</th>
                    <th>Role</th>
                    <th>Share of Voice %</th>
                    <th>Avg AI Rank</th>
                    <th>Top 3 Share %</th>
                    <th>Sentiment</th>
                </tr>
            </thead>
            <tbody>
                {brand_rows_html}
            </tbody>
        </table>

        <h3>🎯 2. Top 10 Prioritised Target Domains (AI-Cited Sources)</h3>
        <table>
            <thead>
                <tr>
                    <th>Target Domain</th>
                    <th>Priority Score</th>
                    <th>Citations</th>
                    <th>Pitch Strategy</th>
                    <th>Citation Worthiness</th>
                    <th>Verified Contact</th>
                </tr>
            </thead>
            <tbody>
                {top10_targets_html}
            </tbody>
        </table>

        <h3>⚔️ 3. Top High-Impact Competitor Gap Pages (Immediate Pitch Targets)</h3>
        <table>
            <thead>
                <tr>
                    <th>Target Page Title & URL</th>
                    <th>Domain</th>
                    <th>Competitors Featured</th>
                    <th>Pitch Strategy</th>
                    <th>Contact Info</th>
                </tr>
            </thead>
            <tbody>
                {top5_gaps_html}
            </tbody>
        </table>

        <h3>💡 4. Top 5 Recommended Linkable Asset & Guest Post Ideas</h3>
        <p style="color:#64748B; font-size:0.92rem; margin-bottom:14px;">
            Content formats engineered to answer buyer evaluation queries and earn direct citations from Google Gemini:
        </p>
        {ideas_html}

        <div class="footer">
            Generated by <strong>{agency_name}</strong> | Confidential Client Deliverable | Powered by AI Citation Link Prospector
        </div>
    </div>
</body>
</html>
"""
    return html_content
